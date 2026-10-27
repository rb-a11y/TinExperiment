"""Load peak lists, MFP candidates and the isotope table from input files.

The scientist's workbook keeps several tables side by side on one sheet, so
the loader finds each table by its header text rather than by fixed cell
positions:

* "Mass (u)" starts a peak list (m/z, then intensity in the next column).
  The title above it says whether it is normalised or raw.
* "dev [ppm]" starts the MFP output. The column to its left is m/z, then
  come DBE, intensity ("Count") and one count column per element.
* "Mass" next to "Abundance" starts the isotope table, labelled e.g. Sn(120).

Plain peak lists (CSV or a two-column sheet) are also accepted.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

from .isotopes import Isotope

ELEMENT_RE = re.compile(r"^[A-Z][a-z]?$")
ISOTOPE_LABEL_RE = re.compile(r"^\s*([A-Z][a-z]?)\s*\(\s*(\d+)\s*\)\s*$|^\s*(\d+)\s*([A-Z][a-z]?)\s*$")
POLARITY_RE = re.compile(r"\((\+|-)\)")


@dataclass
class Dataset:
    """Everything loaded from one input file."""

    file_name: str
    sample: str
    polarity: str  # "+", "-" or "?"
    peaks: dict[str, pd.DataFrame]  # dataset name -> columns mz, intensity (sorted by mz)
    mfp: pd.DataFrame | None = None  # one row per candidate formula
    isotopes: dict[str, Isotope] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def load(source: str | Path | bytes, file_name: str | None = None) -> Dataset:
    """Load a workbook (.xlsx) or a plain peak list (.csv/.txt)."""
    if isinstance(source, (str, Path)):
        file_name = file_name or Path(source).name
        data = Path(source).read_bytes()
    else:
        data = source
        file_name = file_name or "upload"
    suffix = Path(file_name).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _load_workbook(data, file_name)
    if suffix in (".csv", ".txt", ".tsv"):
        return _load_peak_table(pd.read_csv(io.BytesIO(data), sep=None, engine="python"), file_name)
    raise ValueError(f"Unsupported file type '{suffix}'. Use .xlsx, .csv or .txt.")


# ---------------------------------------------------------------- workbook


def _load_workbook(data: bytes, file_name: str) -> Dataset:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = [[list(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    wb.close()

    peaks: dict[str, pd.DataFrame] = {}
    mfp = None
    isotopes: dict[str, Isotope] = {}
    sample = ""
    warnings: list[str] = []

    for grid in sheets:
        for (r, c), name, block_sample in _find_peak_blocks(grid):
            df = _read_numeric_columns(grid, r + 1, [c, c + 1], ["mz", "intensity"])
            if df.empty:
                continue
            name = _unique(name, peaks)
            peaks[name] = df.sort_values("mz", ignore_index=True)
            sample = sample or block_sample
        if mfp is None:
            mfp = _find_mfp(grid)
        isotopes.update(_find_isotopes(grid))

    if not peaks:
        # Fall back to treating the first sheet as a plain peak table.
        return _load_peak_table(pd.read_excel(io.BytesIO(data)), file_name)

    if mfp is None:
        warnings.append("No MFP output found (no 'dev [ppm]' header). Formula cross-checks are skipped.")
    if not isotopes:
        warnings.append("No isotope table found in the file. NIST values are used.")

    polarity = _polarity(sample or file_name)
    if polarity == "?":
        warnings.append("Polarity not found in the sample name. Check that positive and negative runs are not mixed.")
    return Dataset(file_name, sample or Path(file_name).stem, polarity, peaks, mfp, isotopes, warnings)


def _cell(grid, r, c):
    if r < len(grid) and c < len(grid[r]):
        return grid[r][c]
    return None


def _find_peak_blocks(grid):
    for r, row in enumerate(grid):
        for c, value in enumerate(row):
            if isinstance(value, str) and value.strip().lower() == "mass (u)":
                yield (r, c), _block_name(grid, r, c), _clean_sample(_cell(grid, r, c + 1))


def _block_name(grid, r, c) -> str:
    """Name a peak block from the nearest title above it."""
    for rr in range(r - 1, -1, -1):
        for cc in range(c, max(c - 4, -1), -1):
            text = _cell(grid, rr, cc)
            if isinstance(text, str):
                low = text.lower()
                if "raw" in low:
                    return "Raw"
                if "norm" in low:
                    return "Normalised"
    return "Peaks"


def _unique(name: str, existing: dict) -> str:
    candidate, n = name, 2
    while candidate in existing:
        candidate, n = f"{name} {n}", n + 1
    return candidate


def _clean_sample(value) -> str:
    if not isinstance(value, str):
        return ""
    text = value.strip().rstrip("*").strip()
    # "1 (NAG+SnF2_..._(+)1)" -> "NAG+SnF2_..._(+)1"
    match = re.match(r"^\d+\s*\((.*)\)$", text)
    return match.group(1) if match else text


def _polarity(text: str) -> str:
    match = POLARITY_RE.search(text)
    return match.group(1) if match else "?"


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and np.isfinite(value)


def _read_numeric_columns(grid, start_row, cols, names) -> pd.DataFrame:
    """Read rows until the first column stops being numeric."""
    rows = []
    for r in range(start_row, len(grid)):
        values = [_cell(grid, r, c) for c in cols]
        if not _is_number(values[0]):
            break
        rows.append([v if _is_number(v) else np.nan for v in values])
    return pd.DataFrame(rows, columns=names, dtype=float)


def _find_mfp(grid) -> pd.DataFrame | None:
    for r, row in enumerate(grid):
        for c, value in enumerate(row):
            if isinstance(value, str) and value.strip().lower() == "dev [ppm]":
                return _read_mfp(grid, r, c)
    return None


def _read_mfp(grid, r, c) -> pd.DataFrame:
    elements = []
    cc = c + 3
    while isinstance(_cell(grid, r, cc), str) and ELEMENT_RE.match(_cell(grid, r, cc).strip()):
        elements.append(_cell(grid, r, cc).strip())
        cc += 1
    cols = [c - 1, c, c + 1, c + 2] + list(range(c + 3, c + 3 + len(elements)))
    df = _read_numeric_columns(grid, r + 1, cols, ["mz", "ppm", "dbe", "intensity", *elements])
    for el in elements:
        df[el] = df[el].fillna(0).astype(int)
    df["formula"] = [hill_formula(counts) for counts in df[elements].to_dict("records")]
    df["has_sn"] = df["Sn"] > 0 if "Sn" in df else False
    return df


def hill_formula(counts: dict[str, int]) -> str:
    """Write a formula in Hill order (C, H, then alphabetical), omitting 1s."""
    present = {el: n for el, n in counts.items() if n}
    if "C" in present:
        order = ["C", "H"] + sorted(el for el in present if el not in ("C", "H"))
    else:
        order = sorted(present)
    return "".join(f"{el}{present[el] if present[el] != 1 else ''}" for el in order if el in present)


def _find_isotopes(grid) -> dict[str, Isotope]:
    found = {}
    for r, row in enumerate(grid):
        for c, value in enumerate(row):
            right = _cell(grid, r, c + 1)
            if (
                isinstance(value, str)
                and value.strip().lower() == "mass"
                and isinstance(right, str)
                and right.strip().lower().startswith("abundance")
            ):
                for rr in range(r + 1, len(grid)):
                    label, mass, abundance = (_cell(grid, rr, c - 1), _cell(grid, rr, c), _cell(grid, rr, c + 1))
                    if not (_is_number(mass) and _is_number(abundance)):
                        break
                    iso = _isotope_label(label)
                    if iso:
                        found[iso] = Isotope(iso, float(mass), float(abundance), source="provided table")
    return found


def _isotope_label(label) -> str | None:
    if not isinstance(label, str):
        return None
    match = ISOTOPE_LABEL_RE.match(label)
    if not match:
        return None
    element, number = (match.group(1), match.group(2)) if match.group(1) else (match.group(4), match.group(3))
    return f"{number}{element}"


# -------------------------------------------------------------- peak table


def _load_peak_table(table: pd.DataFrame, file_name: str) -> Dataset:
    """A plain table: pick m/z and intensity columns by name, else the first two numeric ones."""
    lower = {str(col).lower(): col for col in table.columns}
    mz_col = next((lower[k] for k in lower if k in ("mz", "m/z", "mass", "mass (u)")), None)
    int_col = next((lower[k] for k in lower if "intens" in k or k in ("counts", "count", "area")), None)
    numeric = [col for col in table.columns if pd.api.types.is_numeric_dtype(table[col])]
    if mz_col is None or int_col is None:
        if len(numeric) < 2:
            raise ValueError("Could not find an m/z column and an intensity column in the file.")
        mz_col, int_col = numeric[0], numeric[1]
    df = pd.DataFrame({"mz": pd.to_numeric(table[mz_col], errors="coerce"),
                       "intensity": pd.to_numeric(table[int_col], errors="coerce")}).dropna()
    df = df.sort_values("mz", ignore_index=True)
    polarity = _polarity(file_name)
    warnings = ["Plain peak list: no MFP output or isotope table, so NIST values are used and formula checks are skipped."]
    if polarity == "?":
        warnings.append("Polarity not found in the file name. Check that positive and negative runs are not mixed.")
    return Dataset(file_name, Path(file_name).stem, polarity, {"Peaks": df}, None, {}, warnings)
